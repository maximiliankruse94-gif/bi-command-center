"""Änderungserkennung: neue Tabs, geänderte Werte, gelöschte Zeilen, unbekannte Tabs, Duplikate.
Läuft NACH test_data_integrity (Dateinamen-Reihenfolge) auf Kopien der Originaldateien."""
import datetime as dt

import openpyxl
import pytest

from conftest import TMP, needs_data

pytestmark = needs_data


def _edit(name, fn):
    p = TMP / "data" / f"{name}.xlsx"
    wb = openpyxl.load_workbook(p, data_only=True)
    fn(wb)
    wb.save(p)


def _sync():
    from bi.sync import run_sync
    return run_sync("test", "local", force=True)


def test_changed_value_is_updated(synced):
    from bi import db

    def change(wb):
        ws = wb["Oktober 26"]
        col = next(c.column for c in ws[1] if c.value == "Umsatz Gesamt")
        ws.cell(row=6, column=col).value = 99999.99      # 05.10.26
    _edit("daily", change)
    res = _sync()
    assert res["datasets"]["daily"]["aktualisiert"] == 1
    d = db.read_table("daily").set_index("date")
    assert d.loc[dt.date(2026, 10, 5), "umsatz_gesamt"] == pytest.approx(99999.99)


def test_new_quarter_tab_and_unknown_tab():
    def add(wb):
        src = wb["Mails 2026 Q4"]
        ws = wb.create_sheet("Mails 2027 Q1")
        for i, row in enumerate(src.iter_rows(min_row=1, max_row=6, values_only=True), start=1):
            row = list(row)
            if i > 1 and isinstance(row[0], (dt.date, dt.datetime, str)):
                row[0] = dt.datetime(2027, 1, i)
                row[5] = f"NL{i:02d}0127"
            for j, v in enumerate(row, start=1):
                ws.cell(row=i, column=j).value = v
        junk = wb.create_sheet("Notizen Marketing")
        for i in range(1, 15):
            junk.cell(row=i, column=1).value = f"Idee {i}"
            junk.cell(row=i, column=2).value = i
    _edit("newsletter", add)
    res = _sync()
    assert res["datasets"]["newsletter"]["neu"] == 5
    assert any(u["tab"] == "Notizen Marketing" for u in res["unknown_tabs"])
    assert res["status"] == "warnung"
    assert any("Mails 2027 Q1" in t for t in res["new_tabs"])


def test_deleted_row_soft_deleted_and_nothing_else_lost():
    from bi import db
    before = len(db.read_table("newsletter"))

    def delete(wb):
        wb["Mails 2027 Q1"].delete_rows(3)
    _edit("newsletter", delete)
    res = _sync()
    assert res["datasets"]["newsletter"]["entfernt"] == 1
    assert len(db.read_table("newsletter")) == before - 1
    assert len(db.read_table("newsletter", include_deleted=True)) == before  # physisch erhalten


def test_duplicate_row_counted_once():
    from bi import db
    before = len(db.read_table("newsletter"))

    def dup(wb):
        ws = wb["Mails 2027 Q1"]
        vals = [c.value for c in ws[2]]
        r = ws.max_row + 1
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j).value = v
    _edit("newsletter", dup)
    res = _sync()
    assert len(db.read_table("newsletter")) == before
    from sqlalchemy import select
    with db.engine().connect() as c:
        msgs = [r.message for r in c.execute(select(db.dq_issue).where(db.dq_issue.c.run_id == res["run_id"],
                                                                         db.dq_issue.c.category == "duplikat"))]
    assert any("doppelt erfasst" in m for m in msgs)


def test_broken_workbook_keeps_existing_data():
    from bi import db
    before = len(db.read_table("weekly"))
    p = TMP / "data" / "weekly.xlsx"
    backup = p.read_bytes()
    p.write_bytes(b"kaputt")
    try:
        res = _sync()
        assert any("Abruf fehlgeschlagen" in e for e in res["errors"])
        assert len(db.read_table("weekly")) == before
    finally:
        p.write_bytes(backup)


def test_manual_mapping_for_unknown_tab():
    from bi import db
    db.save_manual_mapping("newsletter", "Notizen Marketing", "ignore", "test")
    res = _sync()
    assert not any(u["tab"] == "Notizen Marketing" for u in res["unknown_tabs"])
